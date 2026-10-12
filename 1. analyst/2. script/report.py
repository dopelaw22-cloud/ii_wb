#!/usr/bin/env python3
"""Полный отчёт: Markdown + Excel в «4. reports/», опционально в Telegram.

Печатает пути созданных файлов — бот подхватывает их из stdout и отправляет.

Запуск:  python report.py [--days 30] [--telegram] [--quiet]
"""
import sys, pathlib, argparse, datetime, os
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import wb_lib as L
import views as V
import analyze


def md_report(days):
    now = datetime.datetime.now()
    parts = [
        "# Отчёт по магазину «Вытяжки скрытого монтажа»",
        "",
        "ИП Пермин Е. В. · период %d дней · собран %s" % (days, now.strftime("%d.%m.%Y %H:%M")),
        "",
        "## Сводка", "", "```", V.summary(days), "```", "",
        "## Юнит-экономика", "", "```", V.profit_table(), "```", "",
        "## Реклама", "", "```", V.ads_block(days), "```", "",
        "## Воронка продаж", "", "```", V.funnel_block(days), "```", "",
        "## Остатки", "", "```", V.stock_block(days), "```", "",
        "## Карточки", "", "```", V.cards_block(), "```", "",
        "## Рекомендации", "", V.recs_block(limit=200).replace("РЕКОМЕНДАЦИИ\n\n", ""), "",
        "## Свежесть данных", "", "```", V.freshness(), "```", "",
        "## Лимиты WB API", "", "```", V.limits_block(), "```", "",
    ]
    return "\n".join(parts)


def xlsx_report(path, days):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F5496")

    def sheet(title, headers, rows, widths=None):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.font = head_font
            c.fill = head_fill
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths or [], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        return ws

    wb.remove(wb.active)

    econ = L.qd("""
        SELECT p.vendor_code "артикул", p.nm_id "nmID", u.basis "основание",
               u.price_seller "цена", u.commission_pct "кВВ %%", u.commission_rub "комиссия",
               u.logistics_rub "логистика", u.storage_rub "хранение", u.advert_rub "реклама",
               u.cost_rub "себестоимость", u.tax_rub "налог", u.profit_rub "прибыль",
               u.margin_pct "маржа %%", u.roi_pct "ROI %%", u.buyout_pct "выкуп %%",
               u.volume_l "объём, л"
        FROM wb.unit_econ u JOIN wb.product p USING (nm_id)
        WHERE u.calc_date=(SELECT max(calc_date) FROM wb.unit_econ)
        ORDER BY u.basis, u.profit_rub DESC NULLS LAST""")
    if econ:
        sheet("Юнит-экономика", list(econ[0].keys()),
              [[float(v) if isinstance(v, (int, float)) else v for v in r.values()] for r in econ],
              [16, 12, 11, 10, 8, 10, 11, 10, 9, 14, 9, 10, 9, 8, 9, 9])

    ads = L.qd("""
        SELECT a.name "кампания", a.advert_id "ID", a.status "статус",
               sum(s.sum_rub) "затраты", sum(s.views) "показы", sum(s.clicks) "клики",
               sum(s.orders) "заказы", sum(s.sum_price) "выручка"
        FROM wb.advert a JOIN wb.advert_stat s USING (advert_id)
        WHERE s.nm_id=0 AND s.stat_date >= current_date - %s::int
        GROUP BY a.name, a.advert_id, a.status ORDER BY sum(s.sum_rub) DESC""", (days,))
    if ads:
        sheet("Реклама", list(ads[0].keys()) + ["ДРР %", "CTR %", "CPC"],
              [list(r.values()) + [
                  round(L.num(r["затраты"]) / L.num(r["выручка"]) * 100, 1) if L.num(r["выручка"]) else None,
                  round(L.num(r["клики"]) / L.num(r["показы"]) * 100, 2) if L.num(r["показы"]) else None,
                  round(L.num(r["затраты"]) / L.num(r["клики"]), 1) if L.num(r["клики"]) else None]
               for r in ads], [34, 12, 9, 11, 11, 9, 9, 12, 8, 8, 8])

    fun = L.qd("""
        SELECT p.vendor_code "артикул", sum(f.open_count) "переходы",
               sum(f.cart_count) "в корзину", sum(f.order_count) "заказы",
               sum(f.buyout_count) "выкупы", sum(f.order_sum) "сумма заказов",
               max(f.conv_cart_pct) "конв. в корзину %%", max(f.conv_buyout_pct) "выкуп %%",
               max(f.feedback_rating) "оценка"
        FROM wb.funnel f JOIN wb.product p USING (nm_id)
        WHERE (f.period_start, f.period_end) = (SELECT period_start, period_end FROM wb.funnel
                                   ORDER BY period_end DESC, period_start DESC LIMIT 1)
        GROUP BY p.vendor_code ORDER BY sum(f.order_sum) DESC NULLS LAST""")
    if fun:
        sheet("Воронка", list(fun[0].keys()), [list(r.values()) for r in fun],
              [16, 11, 12, 10, 10, 14, 17, 10, 9])

    recs = L.qd("""SELECT severity "важность", scope "раздел", title "находка",
                          detail "что делать"
                   FROM wb.recommendation WHERE status='new'
                   ORDER BY CASE severity WHEN 'critical' THEN 0 WHEN 'warn' THEN 1 ELSE 2 END, id""")
    if recs:
        ws = sheet("Рекомендации", list(recs[0].keys()), [list(r.values()) for r in recs],
                   [11, 10, 52, 95])
        colors = {"critical": "FFC7CE", "warn": "FFEB9C", "info": "DDEBF7"}
        for i, r in enumerate(recs, start=2):
            fill = PatternFill("solid", fgColor=colors.get(r["важность"], "FFFFFF"))
            ws.cell(row=i, column=1).fill = fill
            ws.cell(row=i, column=4).alignment = Alignment(wrap_text=True, vertical="top")

    cards = L.qd("""SELECT vendor_code "артикул", nm_id "nmID", title "название",
                           brand "бренд", photos "фото", chars "характеристики",
                           desc_len "длина описания",
                           round((width*height*length)/1000.0, 2) "объём, л",
                           weight_brutto "вес, кг"
                    FROM wb.product ORDER BY vendor_code""")
    if cards:
        sheet("Карточки", list(cards[0].keys()), [list(r.values()) for r in cards],
              [16, 12, 44, 12, 7, 15, 15, 9, 9])

    if not wb.sheetnames:
        return None
    wb.save(path)
    return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--days", type=int, default=30)
    ap.add_argument("--telegram", action="store_true", help="отправить в Telegram владельцу")
    ap.add_argument("--quiet", action="store_true", help="печатать только пути файлов")
    a = ap.parse_args()

    rid = L.run_start("report.py")
    try:
        analyze.run(a.days, save=True)

        L.REPORTS.mkdir(parents=True, exist_ok=True)
        stamp = datetime.datetime.now().strftime("%Y-%m-%d_%H%M")
        md_path = L.REPORTS / ("отчёт_%s.md" % stamp)
        xl_path = L.REPORTS / ("юнит-экономика_%s.xlsx" % stamp)

        md_path.write_text(md_report(a.days), encoding="utf-8")
        xl = xlsx_report(str(xl_path), a.days)

        if not a.quiet:
            print(V.summary(a.days))
            print()
        print(md_path)
        if xl:
            print(xl_path)

        if a.telegram:
            import bot
            own = bot.owner()
            if not own:
                print("владелец не привязан — напишите боту /start", file=sys.stderr)
            else:
                bot.send(own, V.summary(a.days))
                bot.send(own, V.recs_block(limit=8))
                bot.send_file(own, str(md_path), "полный отчёт")
                if xl:
                    bot.send_file(own, str(xl_path), "юнит-экономика в Excel")
        L.run_end(rid, "ok", rows=1)
    except Exception as e:
        L.run_end(rid, "error", message=str(e)[:400])
        raise


if __name__ == "__main__":
    main()
